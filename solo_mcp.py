#!/usr/bin/env python3
"""Solo-MCP: local, approval-gated GitHub Issues bridge (standard library only).

This is a GitHub-Issues transport MVP, not an MCP protocol server.
"""
import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile

CONFIG_PATH = Path.home() / '.config' / 'solo-mcp' / 'config.json'
SERVICE_FILE = Path.home() / '.config' / 'systemd' / 'user' / 'solo-mcp.service'
SERVICE_NAME = 'solo-mcp.service'
MARKER = 'SOLO_MCP_TASK_V1'
RESULT_MARKER = 'SOLO_MCP_RESULT_V1'
READ_LIMIT = 20000
OUTPUT_LIMIT = 20000
BLOCKED_NAMES = {'.env', '.npmrc', '.pypirc', 'id_rsa', 'id_ed25519', 'credentials', 'credentials.json', 'secrets.json', 'known_hosts'}
BLOCKED_PARTS = {'.ssh', '.gnupg', '.aws', '.git', '.config', '.kube', 'node_modules'}
MUTATING_ACTIONS = frozenset({'run', 'write', 'open_app', 'open_document', 'web_search', 'doc_create', 'doc_append', 'venv_create'})
ALLOWED_TOOLS = {'python', 'python3', 'pytest', 'git', 'javap', 'java', 'jar', 'unzip', 'zipinfo', 'ls', 'find', 'grep', 'rg', 'head', 'tail', 'file', 'stat', 'strings', 'jadx', 'usb-devices', 'lsusb'}


def save_config(config):
    CONFIG_PATH.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    fd, temp = tempfile.mkstemp(prefix='.config-', dir=CONFIG_PATH.parent)
    try:
        with os.fdopen(fd, 'w', encoding='utf-8') as out:
            json.dump(config, out, ensure_ascii=False, indent=2)
            out.write('\n')
        os.chmod(temp, 0o600)
        os.replace(temp, CONFIG_PATH)
    finally:
        if os.path.exists(temp):
            os.unlink(temp)


def load_config():
    if not CONFIG_PATH.is_file():
        raise RuntimeError('Configuração inexistente. Use: python3 solo_mcp.py init --control DONO/REPO_PRIVADO')
    config = json.loads(CONFIG_PATH.read_text(encoding='utf-8'))
    control = config.get('control_repo', '')
    if not re.fullmatch(r'[A-Za-z0-9_.-]+/[A-Za-z0-9_.-]+', control):
        raise RuntimeError('control_repo inválido')
    return config


def gh_token():
    token = os.environ.get('SOLO_MCP_TOKEN', '').strip()
    if token:
        return token
    try:
        result = subprocess.run(['gh', 'auth', 'token'], text=True, capture_output=True, timeout=8, check=True)
        return result.stdout.strip()
    except (OSError, subprocess.SubprocessError):
        raise RuntimeError('Autentique com gh auth login ou defina SOLO_MCP_TOKEN no ambiente')


class Github:
    def __init__(self, token):
        self.token = token

    def request(self, method, endpoint, data=None):
        if not endpoint.startswith('/repos/'):
            raise ValueError('Somente endpoints /repos/... são permitidos')
        url = 'https://api.github.com' + endpoint
        body = json.dumps(data).encode('utf-8') if data is not None else None
        req = urllib.request.Request(url, data=body, method=method, headers={
            'Authorization': 'Bearer ' + self.token,
            'Accept': 'application/vnd.github+json',
            'X-GitHub-Api-Version': '2022-11-28',
            'User-Agent': 'solo-mcp-bridge/0.1',
            'Content-Type': 'application/json',
        })
        try:
            with urllib.request.urlopen(req, timeout=25) as res:
                data = res.read(1024 * 1024)
                return json.loads(data) if data else None
        except urllib.error.HTTPError as ex:
            msg = ex.read(500).decode('utf-8', errors='replace')
            raise RuntimeError(f'GitHub HTTP {ex.code}: {msg}') from ex


def require_private_repo(gh, repo):
    info = gh.request('GET', '/repos/' + repo)
    if not info.get('private', False):
        raise RuntimeError('O repositório de CONTROLE precisa ser PRIVADO. Não publique tarefas em repositórios públicos.')
    return info


def resolve_project(config, project):
    projects = config.get('projects', {})
    if project not in projects:
        raise ValueError('Projeto não cadastrado: ' + str(project))
    root = Path(projects[project]['path']).expanduser().resolve(strict=True)
    if not root.is_dir():
        raise ValueError('Raiz de projeto não é um diretório')
    return root


def safe_path(root, relative, write=False):
    if not isinstance(relative, str) or not relative or '\x00' in relative:
        raise ValueError('Caminho inválido')
    rel = Path(relative)
    if rel.is_absolute() or '..' in rel.parts:
        raise ValueError('Caminho absoluto ou traversal bloqueado')
    if any(p in BLOCKED_PARTS or p in BLOCKED_NAMES or p.startswith('.env.') or p.endswith(('.pem', '.key', '.p12', '.pfx')) for p in rel.parts):
        raise ValueError('Caminho protegido')
    target = (root / rel).resolve(strict=False)
    if not target.is_relative_to(root):
        raise ValueError('Caminho fora da raiz do projeto')
    if write and target.is_symlink():
        raise ValueError('Symlink bloqueado para escrita')
    return target


def approve(description, handler=None):
    print('\n[APROVAÇÃO LOCAL NECESSÁRIA]')
    print(description[:2500])
    response = (input('Autorizar esta operação? [s/N]: ').strip().casefold()
                if handler is None else ('s' if handler(description) else 'n'))
    approved = response == 's'
    print('Operação autorizada.' if approved else 'Operação cancelada.')
    return approved


def do_task(config, task, approval_callback=None):
    if not isinstance(task, dict) or task.get('version') != 1:
        raise ValueError('Formato de tarefa inválido')
    action = task.get('action')
    root = resolve_project(config, task.get('project'))
    if action == 'list':
        path = safe_path(root, task.get('path', '.'))
        if not path.is_dir():
            raise ValueError('Não é um diretório')
        entries = []
        for p in sorted(path.iterdir(), key=lambda x: x.name.lower())[:100]:
            try:
                safe_path(root, str(p.relative_to(root)))
            except ValueError:
                continue
            entries.append({'name': p.name, 'kind': 'dir' if p.is_dir() else 'file'})
        return {'entries': entries, 'truncated': len(list(path.iterdir())) > 100}
    if action == 'read':
        path = safe_path(root, task.get('path'))
        if not path.is_file():
            raise ValueError('Arquivo não encontrado')
        with path.open('rb') as file:
            blob = file.read(READ_LIMIT + 1)
        if b'\x00' in blob:
            raise ValueError('Arquivo binário: use archive_list ou ferramenta local')
        return {'path': str(path.relative_to(root)), 'content': blob[:READ_LIMIT].decode('utf-8', 'replace'), 'truncated': len(blob) > READ_LIMIT}
    if action == 'archive_list':
        path = safe_path(root, task.get('path'))
        if not zipfile.is_zipfile(path):
            raise ValueError('Não é ZIP/JAR/AAR válido')
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()[:100]
            return {'entries': [{'name': i.filename[:250], 'bytes': i.file_size} for i in infos], 'truncated': len(z.infolist()) > 100}
    if action == 'git_status':
        run = subprocess.run(['git', '-C', str(root), 'status', '--short', '--branch'], text=True, capture_output=True, timeout=20)
        return {'returncode': run.returncode, 'stdout': run.stdout[:OUTPUT_LIMIT], 'stderr': run.stderr[:OUTPUT_LIMIT]}
    if action == 'write':
        content = task.get('content')
        if not isinstance(content, str) or len(content.encode('utf-8')) > 32000:
            raise ValueError('Escrita limitada a 32 KB de texto')
        path = safe_path(root, task.get('path'), write=True)
        if not path.parent.is_dir():
            raise ValueError('Pasta de destino inexistente')
        if path.exists() and not path.is_file():
            raise ValueError('Destino não é arquivo regular')
        current = hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else None
        if current != task.get('expected_sha256', None):
            raise ValueError('expected_sha256 não confere (proteção contra sobrescrita)')
        if not approve(f'Gravar {path} ({len(content.encode("utf-8"))} bytes); SHA anterior: {current}', approval_callback):
            return {'status': 'rejected'}
        # Recheck before write to avoid changed file since approval.
        now = hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else None
        if now != current:
            raise ValueError('Arquivo mudou durante aprovação')
        fd, temp = tempfile.mkstemp(prefix='.solo-mcp-', dir=path.parent)
        try:
            with os.fdopen(fd, 'w', encoding='utf-8') as out:
                out.write(content)
            os.replace(temp, path)
        finally:
            if os.path.exists(temp): os.unlink(temp)
        return {'status': 'written', 'path': str(path.relative_to(root)), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}
    if action == 'run':
        argv = task.get('argv')
        if not isinstance(argv, list) or not 1 <= len(argv) <= 30 or not all(isinstance(x, str) and 0 < len(x) <= 2000 for x in argv):
            raise ValueError('argv inválido')
        if argv[0] not in ALLOWED_TOOLS:
            raise ValueError('Executável não está na lista permitida')
        if not approve(f'Projeto: {root}\nExecutar sem shell: {json.dumps(argv, ensure_ascii=False)}', approval_callback):
            return {'status': 'rejected'}
        seconds = min(90, max(1, int(task.get('timeout', 30))))
        try:
            run = subprocess.run(argv, cwd=root, text=True, capture_output=True, timeout=seconds, env={**os.environ, 'PYTHONUNBUFFERED': '1'})
            return {'returncode': run.returncode, 'stdout': run.stdout[:OUTPUT_LIMIT], 'stderr': run.stderr[:OUTPUT_LIMIT]}
        except subprocess.TimeoutExpired:
            return {'status': 'timeout', 'seconds': seconds}
    if action == 'open_app':
        app = task.get('app')
        commands = {
            'writer': ['libreoffice', '--writer'],
            'files': ['xdg-open', str(root)],
            'browser': ['xdg-open', 'https://www.google.com/'],
        }
        if app not in commands:
            raise ValueError('Aplicativo não permitido; use writer, files ou browser')
        if not approve(f'Abrir aplicativo {app} no ambiente gráfico?', approval_callback):
            return {'status': 'rejected'}
        command = commands[app]
        subprocess.Popen(command, cwd=root, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True)
        return {'status': 'launched', 'app': app}
    if action == 'web_search':
        query = task.get('query')
        if not isinstance(query, str) or not 1 <= len(query.strip()) <= 300:
            raise ValueError('Pesquisa inválida')
        url = 'https://www.google.com/search?' + urllib.parse.urlencode({'q': query})
        if not approve(f'Abrir no navegador pesquisa Google:\n{query}', approval_callback):
            return {'status': 'rejected'}
        subprocess.Popen(['xdg-open', url], cwd=root, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True)
        return {'status': 'launched', 'url': url}
    if action == 'open_document':
        path = safe_path(root, task.get('path'))
        if not path.is_file():
            raise ValueError('Documento inexistente')
        if not approve(f'Abrir arquivo com aplicativo padrão:\n{path}', approval_callback):
            return {'status': 'rejected'}
        subprocess.Popen(['xdg-open', str(path)], cwd=root, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True)
        return {'status': 'launched', 'path': str(path.relative_to(root))}
    if action in ('doc_create', 'doc_append'):
        path = safe_path(root, task.get('path'), write=True)
        text = task.get('text')
        if path.suffix.lower() != '.docx':
            raise ValueError('Somente documentos .docx')
        if not isinstance(text, str) or not 0 < len(text.encode('utf-8')) <= 20000:
            raise ValueError('Texto do documento inválido')
        if not path.parent.is_dir() or path.is_symlink():
            raise ValueError('Destino inválido')
        exists = path.exists()
        if action == 'doc_create' and exists:
            raise ValueError('Documento já existe; utilize doc_append')
        if action == 'doc_append' and (not exists or not path.is_file()):
            raise ValueError('Documento inexistente; utilize doc_create')
        old = hashlib.sha256(path.read_bytes()).hexdigest() if exists else None
        if old != task.get('expected_sha256', None):
            raise ValueError('expected_sha256 não corresponde ao documento')
        if not approve(f'{action}: {path}\nConteúdo ({len(text)} caracteres):\n{text[:1800]}', approval_callback):
            return {'status': 'rejected'}
        now = hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else None
        if now != old:
            raise ValueError('Documento mudou durante aprovação')
        try:
            from docx import Document
        except ImportError as ex:
            raise RuntimeError('Instale python-docx no ambiente Python do aplicativo antes de usar DOCX') from ex
        doc = Document(str(path)) if exists else Document()
        for line in text.splitlines():
            doc.add_paragraph(line)
        fd, temp = tempfile.mkstemp(prefix='.solo-mcp-docx-', suffix='.docx', dir=path.parent)
        os.close(fd)
        try:
            doc.save(temp)
            if not exists and path.exists():
                raise ValueError('Documento foi criado durante a operação')
            if exists and hashlib.sha256(path.read_bytes()).hexdigest() != old:
                raise ValueError('Documento foi alterado antes da gravação')
            os.replace(temp, path)
        finally:
            if os.path.exists(temp):
                os.unlink(temp)
        return {'status': 'written', 'path': str(path.relative_to(root)),
                'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}
    if action == 'venv_create':
        directory = task.get('path', '.venv')
        target = safe_path(root, directory, write=True)
        if target.exists() or not target.parent.is_dir():
            raise ValueError('Diretório do ambiente já existe ou pasta pai inexistente')
        if not approve(f'Criar ambiente virtual Python em {target} (NÃO é sandbox contra malware)', approval_callback):
            return {'status': 'rejected'}
        done = subprocess.run([sys.executable, '-m', 'venv', str(target)], cwd=root,
                              text=True, capture_output=True, timeout=120)
        return {'status': 'created' if done.returncode == 0 else 'failed',
                'returncode': done.returncode, 'stderr': done.stderr[:OUTPUT_LIMIT]}
    raise ValueError('Ação desconhecida: ' + str(action))


def parse_task(body):
    if not isinstance(body, str) or not body.startswith(MARKER + '\n'):
        return None
    try:
        task = json.loads(body.split('\n', 1)[1])
    except json.JSONDecodeError:
        return None
    return task if isinstance(task, dict) and task.get('version') == 1 else None


def process_once(config, gh, background=False, issue_number=None, sensitive_only=False):
    repo = config['control_repo']
    info = require_private_repo(gh, repo)
    owner = info['owner']['login'].lower()
    qs = urllib.parse.urlencode({'state': 'open', 'per_page': 100, 'sort': 'created', 'direction': 'asc'})
    issues = gh.request('GET', '/repos/' + repo + '/issues?' + qs)
    total = 0
    for issue in issues:
        if issue.get('pull_request') or issue.get('user', {}).get('login', '').lower() != owner:
            continue
        task = parse_task(issue.get('body', ''))
        if task is None:
            continue
        number = int(issue['number'])
        if issue_number is not None and number != issue_number:
            continue
        if background and task.get('action') in MUTATING_ACTIONS:
            continue  # sem TTY: deixe pendente para aprovação explícita no CLI
        if sensitive_only and task.get('action') not in MUTATING_ACTIONS:
            continue
        path = '/repos/' + repo + '/issues/' + str(number)
        previous = gh.request('GET', path + '/comments?per_page=100')
        if any(c.get('body', '').startswith(RESULT_MARKER) for c in previous):
            gh.request('PATCH', path, {'state': 'closed'})
            continue
        print(f'[TASK #{number}] {task.get("project")} / {task.get("action")}')
        try:
            result = {'ok': True, 'result': do_task(config, task)}
        except Exception as exc:
            result = {'ok': False, 'error': str(exc)[:800]}
        body = RESULT_MARKER + '\n```json\n' + json.dumps(result, ensure_ascii=False, indent=2)[:50000] + '\n```'
        gh.request('POST', path + '/comments', {'body': body})
        gh.request('PATCH', path, {'state': 'closed'})
        total += 1
    return total


def serve_tasks(config, gh, once=False, interval=10, background=False):
    """Consulta tarefas continuamente; só imprime quando há trabalho ou erro."""
    while True:
        try:
            total = process_once(config, gh, background=background)
            if total:
                print(f'Tarefas executadas: {total}', flush=True)
        except Exception as ex:
            print('Erro na consulta:', ex, file=sys.stderr, flush=True)
            if once:
                raise
        if once:
            return
        time.sleep(max(10, interval))


def _unit_quote(value):
    """Quote absolute command paths for systemd, including %-specifier escaping."""
    value = str(value)
    if not value.startswith('/') or any(c in value for c in '\r\n\x00'):
        raise ValueError('Caminho absoluto inválido para unidade systemd')
    return '"' + value.replace('\\', '\\\\').replace('"', '\\"').replace('%', '%%') + '"'


def _unit_working_directory(value):
    """WorkingDirectory= accepts an unquoted absolute path, not shell arguments."""
    value = str(value)
    if not value.startswith('/') or any(c in value for c in '\r\n\x00"'):
        raise ValueError('Diretório de trabalho inválido para unidade systemd')
    return value.replace('\\', '\\\\').replace('%', '%%')


def generate_user_unit(script_path=None, python_path=None):
    """Generate a background-only, non-root systemd user unit."""
    script = Path(script_path or __file__).resolve()
    python = Path(python_path or sys.executable).resolve()
    return (
        '[Unit]\n'
        'Description=Solo-MCP - ponte de tarefas GitHub\n'
        'Wants=network-online.target\n'
        'After=network-online.target\n\n'
        '[Service]\n'
        'Type=simple\n'
        f'ExecStart={_unit_quote(python)} {_unit_quote(script)} serve --background\n'
        f'WorkingDirectory={_unit_working_directory(script.parent)}\n'
        'Restart=on-failure\n'
        'RestartSec=10\n'
        'NoNewPrivileges=true\n'
        'UMask=0077\n'
        'PrivateTmp=true\n\n'
        '[Install]\n'
        'WantedBy=default.target\n'
    )


def manage_service(action):
    """Install and manage a service in the current user's systemd session."""
    if sys.platform != 'linux':
        raise RuntimeError('Serviço systemd --user disponível somente no Linux')
    base = ['systemctl', '--user']
    if action in ('install', 'repair'):
        load_config()  # Evita instalar serviço sem configuração
        SERVICE_FILE.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
        content = generate_user_unit()
        if SERVICE_FILE.exists():
            current = SERVICE_FILE.read_text(encoding='utf-8')
            if current != content:
                if action != 'repair':
                    raise RuntimeError('Unidade antiga difere da atual. Execute: python3 solo_mcp.py service repair')
                if not current.startswith('[Unit]\nDescription=Solo-MCP - ponte de tarefas GitHub\n'):
                    raise RuntimeError('Unidade existente não parece pertencer ao Solo-MCP. Não será substituída.')
                backup = SERVICE_FILE.with_suffix('.service.bak')
                if backup.exists():
                    raise RuntimeError(f'Backup de unidade já existente: {backup}. Revise antes de reparar.')
                backup.write_text(current, encoding='utf-8')
                backup.chmod(0o600)
                print('Backup da unidade anterior:', backup)
        SERVICE_FILE.write_text(content, encoding='utf-8')
        SERVICE_FILE.chmod(0o600)
        subprocess.run(base + ['daemon-reload'], check=True)
        if action == 'repair':
            subprocess.run(base + ['reset-failed', SERVICE_NAME], check=False)
        subprocess.run(base + ['enable', '--now', SERVICE_NAME], check=True)
        if action == 'repair':
            subprocess.run(base + ['restart', SERVICE_NAME], check=True)
        active = subprocess.run(base + ['is-active', '--quiet', SERVICE_NAME], check=False)
        if active.returncode != 0:
            raise RuntimeError('Unidade instalada, mas não está ativa. Consulte: python3 solo_mcp.py service logs')
        print('Serviço ativo. Execute: python3 solo_mcp.py service status')
        return
    if action == 'uninstall':
        subprocess.run(base + ['disable', '--now', SERVICE_NAME], check=True)
        SERVICE_FILE.unlink(missing_ok=True)
        subprocess.run(base + ['daemon-reload'], check=True)
        print('Serviço removido; seus projetos e configuração foram preservados.')
        return
    if action == 'logs':
        subprocess.run(['journalctl', '--user', '-u', SERVICE_NAME,
                        '-n', '50', '--no-pager'], check=True)
        return
    if action not in ('start', 'stop', 'restart', 'status'):
        raise ValueError('Ação de serviço não suportada: ' + action)
    argv = base + [action]
    if action == 'status':
        argv += ['--no-pager', SERVICE_NAME]
        # systemctl status pode retornar 3 se o serviço estiver parado
        subprocess.run(argv, check=False)
    else:
        subprocess.run(argv + [SERVICE_NAME], check=True)


def pending_tasks(config, gh):
    """Show mutating tasks awaiting user confirmation, without exposing content."""
    repo = config['control_repo']
    info = require_private_repo(gh, repo)
    owner = info['owner']['login'].lower()
    query = urllib.parse.urlencode({'state': 'open', 'per_page': 100})
    issues = gh.request('GET', '/repos/' + repo + '/issues?' + query)
    pending = []
    for issue in issues:
        if issue.get('pull_request') or issue.get('user', {}).get('login', '').lower() != owner:
            continue
        task = parse_task(issue.get('body', ''))
        if task and task.get('action') in MUTATING_ACTIONS:
            pending.append((int(issue['number']), str(task.get('project')), task['action']))
    return sorted(pending)


def cli():
    ap = argparse.ArgumentParser(description='Solo-MCP: ponte local via GitHub Issues privadas')
    sub = ap.add_subparsers(dest='command', required=True)
    init = sub.add_parser('init', help='Configurar repositório de controle PRIVADO')
    init.add_argument('--control', required=True)
    add = sub.add_parser('add-project', help='Cadastrar projeto local')
    add.add_argument('--name', required=True)
    add.add_argument('--path', required=True)
    add.add_argument('--repo', default='')
    run = sub.add_parser('serve', help='Escutar novas tarefas')
    run.add_argument('--once', action='store_true')
    run.add_argument('--interval', type=int, default=10)
    run.add_argument('--background', action='store_true', help='Sem TTY: deixar run/write pendentes de aprovação')
    service = sub.add_parser('service', help='Gerenciar serviço systemd --user')
    service.add_argument('action', choices=['install', 'repair', 'start', 'stop', 'restart', 'status', 'logs', 'uninstall'])
    sub.add_parser('pending', help='Mostrar tarefas aguardando aprovação')
    sub.add_parser('desktop', help='Abrir painel persistente de aprovações no ambiente gráfico')
    desktop_autostart = sub.add_parser('desktop-autostart', help='Gerenciar início do painel na sessão gráfica')
    desktop_autostart.add_argument('action', choices=['install', 'uninstall'])
    approve_cmd = sub.add_parser('approve', help='Processar tarefa mutável com confirmação S/N no terminal')
    approve_cmd.add_argument('--issue', type=int, required=True)
    args = ap.parse_args()
    if args.command == 'init':
        if CONFIG_PATH.exists():
            ap.error(f'Configuração já existe: {CONFIG_PATH}')
        save_config({'control_repo': args.control, 'projects': {}})
        print('Configuração local criada:', CONFIG_PATH)
    elif args.command == 'add-project':
        conf = load_config()
        root = Path(args.path).expanduser().resolve(strict=True)
        if not root.is_dir():
            ap.error('Caminho não é diretório')
        conf.setdefault('projects', {})[args.name] = {'path': str(root), 'github': args.repo}
        save_config(conf)
        print(f'Projeto {args.name} registrado: {root}')
    elif args.command == 'service':
        manage_service(args.action)
    elif args.command == 'desktop':
        from solo_desktop import start_desktop
        start_desktop()
    elif args.command == 'desktop-autostart':
        from solo_desktop import manage_autostart
        manage_autostart(args.action)
    elif args.command == 'pending':
        conf = load_config()
        gh = Github(gh_token())
        tasks = pending_tasks(conf, gh)
        if not tasks:
            print('Nenhuma tarefa aguardando aprovação.')
        for number, project, action in tasks:
            print(f'#{number} projeto={project} ação={action} | python3 solo_mcp.py approve --issue {number}')
    elif args.command == 'approve':
        if not sys.stdin.isatty():
            ap.error('Aprovação exige terminal interativo (TTY)')
        conf = load_config()
        gh = Github(gh_token())
        count = process_once(conf, gh, issue_number=args.issue, sensitive_only=True)
        if count == 0:
            print('Issue não encontrada, já concluída ou não requer aprovação.')
    elif args.command == 'serve':
        conf = load_config()
        gh = Github(gh_token())
        require_private_repo(gh, conf['control_repo'])
        print('Ponte conectada a', conf['control_repo'])
        if not args.once:
            mode = 'somente leitura' if args.background else 'interativo'
            print(f'Aguardando tarefas ({mode}, consulta a cada {max(10, args.interval)}s). Ctrl+C para sair.')
        try:
            serve_tasks(conf, gh, once=args.once, interval=args.interval, background=args.background)
        except KeyboardInterrupt:
            print('\nPonte encerrada.')


if __name__ == '__main__':
    cli()
